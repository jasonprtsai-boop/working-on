import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

from backend.utils.serialization.excel_exporter import ExcelExporter


class ExcelExporterReportTest(unittest.TestCase):
    def test_report_adds_action_items_and_filters_snapshot_noise(self) -> None:
        fen = "rnbakabnr/9/1c5c1/p1p1p1p1p/9/9/P1P1P1P1P/1C5C1/9/RNBAKABNR w - - 0 1"
        events = [
            {
                "sequence_id": 1,
                "event_id": "11111111-1111-4111-8111-111111111111",
                "session_id": "session-a",
                "trace_id": "trace-a",
                "timestamp": 1788766600.0,
                "type": "STATE_UPDATED",
                "source": "state_manager",
                "payload": {
                    "vision": {
                        "detections_count": 0,
                        "avg_confidence": 0.0,
                        "min_confidence": 0.0,
                        "latency_ms": 5,
                        "fen": fen,
                    },
                    "engine": {"bestmove": "b0c2", "score": 0.3, "depth": 8},
                    "robot": {"connected": False, "error": "snapshot only"},
                },
            },
            {
                "sequence_id": 2,
                "event_id": "22222222-2222-4222-8222-222222222222",
                "session_id": "session-a",
                "trace_id": "trace-a",
                "timestamp": 1788766601.0,
                "type": "VISION_FRAME_PROCESSED",
                "source": "vision_service",
                "payload": {
                    "detections_count": 1,
                    "avg_confidence": 0.91,
                    "min_confidence": 0.91,
                    "latency_ms": 42,
                    "camera_status": "ready",
                    "fen": fen,
                },
            },
            {
                "sequence_id": 3,
                "event_id": "33333333-3333-4333-8333-333333333333",
                "session_id": "session-a",
                "trace_id": "trace-a",
                "timestamp": 1788766602.0,
                "type": "ROBOT_STATUS_UPDATED",
                "source": "robot_status_worker",
                "payload": {
                    "connected": False,
                    "is_connected": False,
                    "error": "link failed",
                    "ip": "192.168.10.10",
                    "port": 502,
                    "connection": {"adapter": "modbus", "connected": False},
                    "robot_status": "disconnected",
                },
            },
            {
                "sequence_id": 4,
                "event_id": "44444444-4444-4444-8444-444444444444",
                "session_id": "session-a",
                "trace_id": "trace-a",
                "timestamp": 1788766603.0,
                "type": "ENGINE_INFO_UPDATED",
                "source": "AI_ENGINE",
                "payload": {"bestmove": "h2e2", "score": 12, "depth": 10, "engine_ms": 24},
            },
        ]

        with tempfile.TemporaryDirectory() as tmp:
            output = Path(tmp) / "report.xlsx"
            exporter = ExcelExporter(filename=str(Path(tmp) / "runtime.xlsx"), subscribe=False)
            exporter.export_events(events, str(output), session_id="session-a")

            wb = load_workbook(output, data_only=True)
            self.addCleanup(wb.close)

            self.assertIn("Action Items", wb.sheetnames)
            overview = {
                row[0]: row[1]
                for row in wb["Overview"].iter_rows(min_row=4, max_col=2, values_only=True)
                if row[0]
            }
            self.assertEqual(overview["Report Verdict"], "NOT READY - robot disconnected")
            self.assertEqual(overview["Game Moves"], 0)
            self.assertEqual(overview["Vision Events"], 1)
            self.assertEqual(overview["YOLO Min Confidence"], 0.91)

            action_rows = list(wb["Action Items"].iter_rows(min_row=2, values_only=True))
            self.assertEqual(action_rows[0][0], "P0")
            self.assertEqual(action_rows[0][1], "Robot")
            self.assertIn("endpoint=192.168.10.10:502", action_rows[0][3])
            self.assertIn("adapter=modbus", action_rows[0][3])
            self.assertIn("error=link failed", action_rows[0][3])
            self.assertEqual([row[1] for row in action_rows].count("Robot"), 1)

            self.assertEqual(wb["Vision YOLO"].max_row, 2)
            self.assertEqual(wb["Game Moves"].max_row, 1)
            self.assertEqual(wb["Engine AI"].max_row, 2)
            self.assertEqual(wb["UCCI Trace"].max_row, 2)

            warnings = list(wb["Errors & Warnings"].iter_rows(min_row=2, values_only=True))
            reasons = [row[1] for row in warnings]
            self.assertNotIn("Low YOLO confidence", reasons)
            self.assertIn("link failed", " ".join(str(row[1]) for row in warnings))

    def test_field_profile_focuses_report_and_excludes_raw_sheets(self) -> None:
        events = [
            {
                "sequence_id": 1,
                "event_id": "55555555-5555-4555-8555-555555555555",
                "session_id": "session-field",
                "trace_id": "trace-field",
                "timestamp": 1788766700.0,
                "type": "GAME_PLAYER_MOVE",
                "source": "player",
                "payload": {"move": "a0a1", "fen_after": "fen-after"},
            },
            {
                "sequence_id": 2,
                "event_id": "66666666-6666-4666-8666-666666666666",
                "session_id": "session-field",
                "trace_id": "trace-field",
                "timestamp": 1788766701.0,
                "type": "ROBOT_STATUS_UPDATED",
                "source": "robot",
                "payload": {"connected": True, "robot_status": "ready", "message": "field test ok"},
            },
        ]

        with tempfile.TemporaryDirectory() as tmp:
            output = Path(tmp) / "field-report.xlsx"
            exporter = ExcelExporter(filename=str(Path(tmp) / "runtime.xlsx"), subscribe=False)
            exporter.export_events(events, str(output), session_id="session-field", profile="field")

            wb = load_workbook(output, data_only=True)
            self.addCleanup(wb.close)

            self.assertEqual(wb["Overview"]["A1"].value, "S.M.A.R.T Chess Robot Field Report")
            self.assertIn("Field Test Report", wb.sheetnames)
            self.assertIn("Game Moves", wb.sheetnames)
            self.assertIn("Errors & Warnings", wb.sheetnames)
            self.assertNotIn("Pipeline_Log", wb.sheetnames)
            self.assertNotIn("Raw Payload", wb.sheetnames)
            self.assertNotIn("Vision Detections", wb.sheetnames)


if __name__ == "__main__":
    unittest.main()
